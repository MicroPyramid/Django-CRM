import datetime
import io
import re
import time
import zipfile

import openpyxl
import pytest
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import RequestFactory

from common.spreadsheet import (
    _XLSX_BLANK_RUN,
    MAX_COLUMNS,
    MAX_UPLOAD_BYTES,
    SpreadsheetError,
    Table,
    read_table,
    read_upload,
)

NOT_OPENED = "That .xlsx file could not be opened. Save it again, or export it as .csv."
TOO_LARGE = "That .xlsx file is too large to read."


def _xlsx(rows, extra_sheet=None):
    wb = openpyxl.Workbook()
    ws = wb.active
    for row in rows:
        ws.append(row)
    if extra_sheet is not None:
        ws2 = wb.create_sheet()
        for row in extra_sheet:
            ws2.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _zip(members):
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zf:
        for name, content in members:
            zf.writestr(name, content)
    return buf.getvalue()


# CSV


def test_csv_utf8_with_bom():
    data = "﻿Name,Email\nAda,a@x.com\n".encode()
    assert read_table("a.csv", data, 10) == Table(
        ["Name", "Email"], [["Ada", "a@x.com"]]
    )


def test_csv_windows_1252():
    table = read_table("a.csv", "Name\nJosé\n".encode("cp1252"), 10)
    assert table.rows == [["José"]]


def test_csv_semicolon_and_tab():
    for text in ("a;b\n1;2\n", "a\tb\n1\t2\n"):
        table = read_table("a.csv", text.encode(), 10)
        assert table.headers == ["a", "b"]
        assert table.rows == [["1", "2"]]


def test_csv_leading_blank_rows_skipped():
    table = read_table("a.csv", b",,\nName,Email\nAda,a@x.com\n", 10)
    assert table.headers == ["Name", "Email"]
    assert table.rows == [["Ada", "a@x.com"]]


def test_csv_blank_rows_dropped():
    table = read_table("a.csv", b"Name,Email\nAda,a@x.com\n\n,\nBob,b@x.com\n", 10)
    assert table.rows == [["Ada", "a@x.com"], ["Bob", "b@x.com"]]


def test_csv_cells_are_stripped():
    table = read_table("a.csv", b" Name , Email\n Ada , a@x.com \n", 10)
    assert table.headers == ["Name", "Email"]
    assert table.rows == [["Ada", "a@x.com"]]


@pytest.mark.parametrize("name", ["a.csv", "a.xlsx"])
def test_leading_apostrophe_before_formula_char_is_stripped(name):
    # The skipped-rows download and the server's exports put one "'" before a
    # cell a spreadsheet would run as a formula; reading the file takes it off.
    rows = [
        ["Name", "Phone", "Note"],
        ["Ada", "'+44 20 7946 0000", "'=x"],
        ["Bob", "'-5", "'@home"],
        ["Cy", "'hello", "''+1"],
    ]
    if name.endswith(".csv"):
        data = "\n".join(",".join(r) for r in rows).encode()
    else:
        data = _xlsx(rows)
    assert read_table(name, data, 10).rows == [
        ["Ada", "+44 20 7946 0000", "=x"],
        ["Bob", "-5", "@home"],
        ["Cy", "'hello", "'+1"],
    ]


def test_csv_nul_refused():
    with pytest.raises(SpreadsheetError) as exc:
        read_table("a.csv", b"Name\nA\x00da\n", 10)
    assert str(exc.value) == "The file could not be read: it contains NUL characters."


def test_csv_field_over_limit_is_a_clean_error():
    big = b"Name\n" + b"a" * 200_000 + b"\n"
    with pytest.raises(SpreadsheetError):
        read_table("a.csv", big, 10)


# XLSX


def test_xlsx_first_sheet_only():
    data = _xlsx([["Name"], ["Ada"]], extra_sheet=[["Other"], ["Zed"]])
    assert read_table("a.xlsx", data, 10) == Table(["Name"], [["Ada"]])


def test_xlsx_values():
    row = [
        919876543210.0,
        42,
        datetime.date(2026, 10, 20),
        datetime.datetime(2026, 10, 20, 0, 0),
        None,
        True,
    ]
    data = _xlsx([["a", "b", "c", "d", "e", "f"], row])
    table = read_table("a.xlsx", data, 10)
    assert table.rows == [
        ["919876543210", "42", "2026-10-20", "2026-10-20", "", "true"]
    ]


def test_xlsx_formula_uses_cached_value():
    data = _xlsx([["a"], ["=1+1"]])
    assert read_table("a.xlsx", data, 10).rows == []
    data = _xlsx([["a", "b"], ["x", "=1+1"]])
    assert read_table("a.xlsx", data, 10).rows == [["x", ""]]


def test_xlsx_not_a_zip():
    with pytest.raises(SpreadsheetError) as exc:
        read_table("a.xlsx", b"not a zip", 10)
    assert str(exc.value) == NOT_OPENED


def test_xlsx_zip_that_is_not_a_workbook():
    with pytest.raises(SpreadsheetError) as exc:
        read_table("a.xlsx", _zip([("hello.txt", "hi")]), 10)
    assert str(exc.value) == NOT_OPENED


def test_xlsx_zip_bomb_refused():
    data = _zip([("big.bin", b"\x00" * (60 * 1024 * 1024))])
    with pytest.raises(SpreadsheetError) as exc:
        read_table("a.xlsx", data, 10)
    assert str(exc.value) == TOO_LARGE


def test_xlsx_too_many_entries():
    data = _zip([(f"f{i}.txt", b"") for i in range(1001)])
    with pytest.raises(SpreadsheetError) as exc:
        read_table("a.xlsx", data, 10)
    assert str(exc.value) == TOO_LARGE


def _with_dimension(data: bytes, ref: str) -> bytes:
    """Rewrite the first sheet's declared size, as a crafted file would."""
    src = zipfile.ZipFile(io.BytesIO(data))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst:
        for info in src.infolist():
            content = src.read(info.filename)
            if info.filename == "xl/worksheets/sheet1.xml":
                content, n = re.subn(
                    rb'<dimension ref="[^"]*"',
                    f'<dimension ref="{ref}"'.encode(),
                    content,
                )
                assert n == 1
            dst.writestr(info.filename, content)
    return out.getvalue()


def _far_corner_cell():
    """A real value at XFD1048576: a million-row gap the reader fills in."""
    wb = openpyxl.Workbook()
    wb.active.append(["Name", "Email"])
    wb.active.append(["Ada", "a@x.com"])
    wb.active.cell(row=1048576, column=16384, value="x")
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


@pytest.mark.parametrize("shape", ["declared", "far_corner"])
def test_xlsx_huge_declared_dimension_is_bounded(shape):
    # The reader pads every row to the declared size: 16,384 cells a row for
    # a million rows, from a file of a few kilobytes.
    if shape == "declared":
        data = _with_dimension(
            _xlsx([["Name", "Email"], ["Ada", "a@x.com"]]), "A1:XFD1048576"
        )
    else:
        data = _far_corner_cell()
    assert len(data) < 10_000
    started = time.monotonic()
    table = read_table("a.xlsx", data, 5000)
    assert time.monotonic() - started < 2
    assert table == Table(["Name", "Email"], [["Ada", "a@x.com"]])


def test_xlsx_wide_declared_dimension_with_few_real_columns_imports():
    # Formatting on an empty cell far to the right widens the declared size
    # without adding a column.
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Name", "Email"])
    ws.append(["Ada", "a@x.com"])
    ws.cell(row=1, column=16384).font = openpyxl.styles.Font(bold=True)
    ws.cell(row=2, column=50).font = openpyxl.styles.Font(bold=True)
    buf = io.BytesIO()
    wb.save(buf)
    with zipfile.ZipFile(io.BytesIO(buf.getvalue())) as zf:
        assert b'<dimension ref="A1:XFD2"' in zf.read("xl/worksheets/sheet1.xml")
    table = read_table("a.xlsx", buf.getvalue(), 10)
    assert table == Table(["Name", "Email"], [["Ada", "a@x.com"]])


def _grid_file(name, grid):
    if name.endswith(".csv"):
        return "".join(",".join(row) + "\n" for row in grid).encode()
    return _xlsx([[cell or None for cell in row] for row in grid])


@pytest.mark.parametrize("name", ["a.csv", "a.xlsx"])
def test_data_under_no_header_is_kept_as_an_unnamed_column(name):
    """A value in a column with no header must reach the matching step, where it
    shows as an unnamed column, rather than vanish uncounted. That includes a
    row whose only value sits there."""
    grid = [
        ["Name", "Email", "", ""],
        ["Ada", "a@x.com", "", "extra"],
        ["", "", "", "only-here"],
    ]
    table = read_table(name, _grid_file(name, grid), 10)
    assert table.headers == ["Name", "Email", "", ""]
    assert table.rows == [["Ada", "a@x.com", "", "extra"], ["", "", "", "only-here"]]


@pytest.mark.parametrize("name", ["a.csv", "a.xlsx"])
def test_trailing_blank_columns_are_trimmed(name):
    grid = [["Name", "Email", "", ""], ["Ada", "a@x.com", "", ""]]
    table = read_table(name, _grid_file(name, grid), 10)
    assert table == Table(["Name", "Email"], [["Ada", "a@x.com"]])


@pytest.mark.parametrize("name", ["a.csv", "a.xlsx"])
@pytest.mark.parametrize("column", [MAX_COLUMNS + 1, MAX_COLUMNS + 2, 1500])
def test_data_past_the_column_limit_is_refused(name, column):
    """Never clipped: a value however far right is either kept or refused."""
    row = [""] * (column - 1) + ["too far"]
    table_file = _grid_file(name, [["Name"], row])
    with pytest.raises(SpreadsheetError) as exc:
        read_table(name, table_file, 10)
    assert str(exc.value) == (
        f"Too many columns (more than {MAX_COLUMNS}); the limit is {MAX_COLUMNS}."
    )


@pytest.mark.parametrize("name", ["a.csv", "a.xlsx"])
def test_too_many_columns_refused(name):
    header = [f"c{i}" for i in range(MAX_COLUMNS + 1)]
    if name.endswith(".csv"):
        data = (",".join(header) + "\n" + "x\n").encode()
    else:
        data = _xlsx([header, ["x"]])
    with pytest.raises(SpreadsheetError) as exc:
        read_table(name, data, 10)
    assert str(exc.value) == (
        f"Too many columns (more than {MAX_COLUMNS}); the limit is {MAX_COLUMNS}."
    )


@pytest.mark.parametrize("name", ["a.csv", "a.xlsx"])
def test_exactly_max_columns_is_allowed(name):
    header = [f"c{i}" for i in range(MAX_COLUMNS)]
    if name.endswith(".csv"):
        data = (",".join(header) + "\n" + "x\n").encode()
    else:
        data = _xlsx([header, ["x"]])
    assert len(read_table(name, data, 10).headers) == MAX_COLUMNS


def test_xlsx_blank_run_ends_the_data():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Name"])
    ws.append(["Ada"])
    # A shorter run of blank rows is still inside the data.
    ws.cell(row=2 + _XLSX_BLANK_RUN, column=1, value="Bob")
    # A full run of blank rows ends it: what follows is not read.
    ws.cell(row=2 + 2 * _XLSX_BLANK_RUN + 1, column=1, value="Cy")
    buf = io.BytesIO()
    wb.save(buf)
    table = read_table("a.xlsx", buf.getvalue(), 5000)
    assert table.rows == [["Ada"], ["Bob"]]


def test_xlsx_rows_read_are_capped_with_blank_rows_counted():
    # Two data rows are within a limit of 2, but 4 * 2 + 1000 rows read in all
    # (blank ones included) are not, even though no blank run ends the data.
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Name"])
    ws.cell(row=600, column=1, value="Ada")
    ws.cell(row=1200, column=1, value="Bob")
    buf = io.BytesIO()
    wb.save(buf)
    with pytest.raises(SpreadsheetError) as exc:
        read_table("a.xlsx", buf.getvalue(), 2)
    assert str(exc.value) == "Too many rows (more than 2); the limit is 2."


def test_xlsx_5000_rows_with_blank_rows_between_imports():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Name", "Email"])
    for i in range(5000):
        ws.append([f"P{i}", f"p{i}@x.com"])
        if i % 10 == 9:
            ws.append([None, None])
    buf = io.BytesIO()
    wb.save(buf)
    table = read_table("a.xlsx", _with_dimension(buf.getvalue(), "A1:XFD5501"), 5000)
    assert len(table.rows) == 5000
    assert table.headers == ["Name", "Email"]
    assert table.rows[-1] == ["P4999", "p4999@x.com"]


def test_defusedxml_active():
    """openpyxl exposes its defusedxml switch as `openpyxl.xml.DEFUSEDXML`."""
    import openpyxl.xml

    assert openpyxl.xml.DEFUSEDXML is True


# Both formats


@pytest.mark.parametrize("name", ["a.xls", "a.ods", "a.txt"])
def test_unsupported_extension(name):
    with pytest.raises(SpreadsheetError) as exc:
        read_table(name, b"x", 10)
    assert str(exc.value) == "Save it as .xlsx or .csv and try again."


def test_extension_is_case_insensitive():
    assert read_table("A.CSV", b"h\nv\n", 10).rows == [["v"]]


@pytest.mark.parametrize("name", ["a.csv", "a.xlsx"])
def test_empty_file(name):
    with pytest.raises(SpreadsheetError) as exc:
        read_table(name, b"", 10)
    assert str(exc.value) == "The file is empty."


def test_header_only_file_has_no_rows():
    assert read_table("a.csv", b"Name,Email\n", 10) == Table(["Name", "Email"], [])


def test_too_many_rows():
    with pytest.raises(SpreadsheetError) as exc:
        read_table("a.csv", b"h\n1\n2\n3\n", 2)
    assert str(exc.value) == "Too many rows (3); the limit is 2."


def test_exactly_max_rows_is_allowed():
    assert len(read_table("a.csv", b"h\n1\n2\n", 2).rows) == 2


def test_too_many_rows_xlsx():
    data = _xlsx([["h"], ["1"], ["2"], ["3"], ["4"]])
    with pytest.raises(SpreadsheetError) as exc:
        read_table("a.xlsx", data, 2)
    assert str(exc.value) == "Too many rows (more than 2); the limit is 2."


def test_exactly_max_rows_is_allowed_xlsx():
    data = _xlsx([["h"], ["1"], ["2"]])
    assert len(read_table("a.xlsx", data, 2).rows) == 2


# read_upload


def test_read_upload_missing_file():
    request = RequestFactory().post("/x/")
    with pytest.raises(SpreadsheetError) as exc:
        read_upload(request, 10)
    assert str(exc.value) == "No file uploaded (expected field 'file')"


def test_read_upload_too_big():
    big = SimpleUploadedFile("a.csv", b"x" * (MAX_UPLOAD_BYTES + 1))
    request = RequestFactory().post("/x/", {"file": big})
    with pytest.raises(SpreadsheetError) as exc:
        read_upload(request, 10)
    assert str(exc.value) == "File exceeds the 5 MB upload limit"


def test_read_upload_ok():
    up = SimpleUploadedFile("a.csv", b"Name\nAda\n")
    request = RequestFactory().post("/x/", {"file": up})
    assert read_upload(request, 10) == Table(["Name"], [["Ada"]])


def _understate_member_size(data: bytes, member: str, declared: int) -> bytes:
    """Patch the declared uncompressed size of one member in both the local
    header and the central directory, leaving the compressed bytes alone."""
    import struct

    out = bytearray(data)
    with zipfile.ZipFile(io.BytesIO(data)) as zf:
        info = zf.getinfo(member)
    real = info.file_size
    # Local header: signature(4) ... compressed size at +18, uncompressed at +22.
    local = info.header_offset
    assert out[local : local + 4] == b"PK\x03\x04"
    assert struct.unpack_from("<I", out, local + 22)[0] == real
    struct.pack_into("<I", out, local + 22, declared)
    # Central directory entry: uncompressed size at +24.
    pos = out.index(b"PK\x01\x02" + b"", 0)
    patched = False
    while pos != -1:
        name_len = struct.unpack_from("<H", out, pos + 28)[0]
        if bytes(out[pos + 46 : pos + 46 + name_len]) == member.encode():
            struct.pack_into("<I", out, pos + 24, declared)
            patched = True
            break
        pos = out.find(b"PK\x01\x02", pos + 4)
    assert patched
    return bytes(out)


def test_xlsx_understated_size_is_refused_not_inflated():
    """_check_zip trusts the declared (attacker-controlled) member sizes.

    That is safe only because zipfile cuts each member's output at its declared
    size and then fails the CRC check, so a member that lies about its size
    cannot be inflated past the bound. This pins that guarantee: a workbook
    whose big part declares a tiny size is refused as unreadable.
    """
    wb = openpyxl.Workbook()
    wb.active.append(["h"])
    # Excel caps a cell at 32,767 characters, so spread the bulk over cells.
    for i in range(10):
        wb.active.append([f"{i}" * 30_000])
    buf = io.BytesIO()
    wb.save(buf)
    with zipfile.ZipFile(io.BytesIO(buf.getvalue())) as zf:
        sizes = {i.filename: i.file_size for i in zf.infolist()}
    member = max(sizes, key=sizes.get)
    assert sizes[member] > 100_000
    data = _understate_member_size(buf.getvalue(), member, 100)
    with pytest.raises(SpreadsheetError) as exc:
        read_table("a.xlsx", data, 10)
    assert str(exc.value) == NOT_OPENED
